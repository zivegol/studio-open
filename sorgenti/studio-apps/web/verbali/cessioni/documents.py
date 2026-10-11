import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from domain import paragraphs


def word(data, path, final=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.5)
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Liberation Serif', Pt(12)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ['Title', 'Heading 1']:
        doc.styles[name].font.name = 'Liberation Serif'
        doc.styles[name].font.size = Pt(15 if name == 'Title' else 12)
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        fmt = doc.styles[name].paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.space_before = Pt(14 if name == 'Heading 1' else 0)
        fmt.space_after = Pt(8 if name == 'Heading 1' else 20)
        fmt.keep_with_next = True
    doc.styles['Heading 1'].font.bold = True
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        for fonts in style.element.iter(qn('w:rFonts')):
            for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
                fonts.attrib.pop(qn('w:'+attr), None)
    if not final:
        p = doc.add_paragraph('BOZZA DA COMPLETARE E VERIFICARE')
        p.runs[0].bold = True
    premise_number = 0
    blocks = [(kind, line) for kind, text in paragraphs(data)
              for line in (text.splitlines() if kind == 'body' else [text]) if line.strip()]
    for kind, text in blocks:
        if kind == 'premise':
            premise_number += 1
            text = f'{premise_number}.\t{text}'
        elif kind == 'share':
            text = '-\t' + text
        p = doc.add_paragraph(text, 'Title' if kind == 'title' else 'Heading 1' if kind == 'heading' else 'Normal')
        if kind in ('premise', 'share'):
            p.paragraph_format.left_indent = Cm(0.65 if kind == 'premise' else 1.25)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.65 if kind == 'premise' else 1.25))
            p.paragraph_format.space_after = Pt(6)
        if kind == 'signature':
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_together = True
        p.paragraph_format.widow_control = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Pagina ').font.size = Pt(9)
    for instruction in ['PAGE', 'NUMPAGES']:
        if instruction == 'NUMPAGES': footer.add_run(' di ').font.size = Pt(9)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), instruction)
        footer._p.append(field)
    doc.core_properties.author = 'studio-example'
    doc.core_properties.title = 'Cessione di quote societarie'
    doc.save(path)


def pdfa(docx, target):
    with tempfile.TemporaryDirectory() as td:
        profile = Path(td) / 'profile'
        subprocess.run(['libreoffice', '-env:UserInstallation='+profile.as_uri(), '--headless', '--convert-to',
                        'pdf:writer_pdf_Export:{"SelectPdfVersion":{"type":"long","value":"1"}}',
                        '--outdir', str(target), str(docx)], check=True, capture_output=True, timeout=90)
    pdf = Path(target) / (Path(docx).stem + '.pdf')
    if not pdf.is_file(): raise ValueError('Conversione PDF/A non riuscita')
    proc = subprocess.run(['/opt/verapdf/verapdf', '--format', 'xml', '--flavour', '1b', str(pdf)], capture_output=True, timeout=90)
    try:
        root = ET.fromstring(proc.stdout)
        reports = [n for n in root.iter() if n.tag.split('}')[-1] == 'validationReport']
        if len(reports) != 1 or reports[0].get('isCompliant') != 'true': raise ValueError()
    except (ET.ParseError, ValueError):
        raise ValueError('PDF/A non conforme: il definitivo non è stato creato')
    return pdf.read_bytes(), proc.stdout.decode('utf-8')
